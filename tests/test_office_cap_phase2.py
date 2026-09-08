# -*- coding: utf-8 -*-
"""阶段二（文档/PPT 能力补齐）回归：S6 pptx_create + S7 markdown→docx 富文本。

覆盖报告《文档表格PPT能力报告_v3.8.5.md》阶段二核心项：
- S6 pptx_create：从 slides 数组或 markdown outline 生成 .pptx（封面/正文/子项
  层级/表格/备注/主题/模板），并完成六层注册 + dc 命名空间导出
- S7 create_doc 的 docx 分支：markdown → Word 富文本（标题样式/加粗/列表/表格/
  代码块/引用/行内），替换原"逐行纯文本段落"
"""
import os

import deepseek_client as dc


def _perm(tmp_path):
    import permissions
    permissions.WORKSPACE_DIR = str(tmp_path)


def _reg(name):
    names = [t["function"]["name"] for t in dc.TOOLS]
    assert name in names, f"{name} 必须注册进 TOOLS"
    assert hasattr(dc, name) and callable(getattr(dc, name)), f"{name} 必须经 dc 导出"
    assert name in dc.TOOL_CALL_MAP, f"{name} 必须在 CALL_MAP"


def test_s6_pptx_create_registered_and_slides(tmp_path):
    _perm(tmp_path)
    _reg("pptx_create")
    from pptx import Presentation
    out = str(tmp_path / "deck.pptx")
    slides = [
        {"title": "项目汇报", "bullets": ["- 完成阶段一", "  - S1 白名单", "- 阶段二"], "notes": "讲3分钟"},
        {"title": "数据", "table": {"headers": ["指标", "值"], "rows": [["营收", "100"], ["毛利", "40"]]}},
    ]
    r = dc.pptx_create(path=out, slides=slides, theme="ocean", title="2026汇报")
    assert r.startswith("已生成"), r
    prs = Presentation(out)
    assert sum(1 for _ in prs.slides) == 3, "封面 + 2 内容页"
    alltxt = ""
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                alltxt += sh.text_frame.text + "\n"
    assert "2026汇报" in alltxt and "项目汇报" in alltxt
    assert "S1 白名单" in alltxt, "子项要点应写入"
    has_tbl = any(getattr(sh, "has_table", False) for s in prs.slides for sh in s.shapes)
    assert has_tbl, "表格应写入"


def test_s6_pptx_create_outline_and_validation(tmp_path):
    _perm(tmp_path)
    from pptx import Presentation
    out = str(tmp_path / "o.pptx")
    r = dc.pptx_create(path=out, outline="# 一\n要点\n\n# 二\n- 甲\n- 乙")
    assert r.startswith("已生成"), r
    assert sum(1 for _ in Presentation(out).slides) == 2, "两个一级标题 = 两页"
    # slides 与 outline 都为空应拒绝；都传也应拒绝
    assert "必填" in dc.pptx_create(path=str(tmp_path / "x.pptx"))
    assert "只能二选一" in dc.pptx_create(path=str(tmp_path / "y.pptx"), slides=[{"title": "a"}], outline="# b")


def test_s7_create_doc_docx_rich_markdown(tmp_path):
    _perm(tmp_path)
    dp = str(tmp_path / "rich.docx")
    md = (
        "# 标题\n\n"
        "这是**加粗**和*斜体*及`代码`。\n\n"
        "- 项目一\n- 项目二\n\n"
        "1. 有序一\n2. 有序二\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n\n"
        "```\nprint('hi')\n```\n\n> 引用行\n"
    )
    r = dc.create_doc(dp, md)
    assert r.startswith("已创建"), r
    from docx import Document
    doc = Document(dp)
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads and "标题" in heads[0], "标题应转为 Word 标题样式"
    assert any(rn.bold for p in doc.paragraphs if "加粗" in p.text for rn in p.runs), "加粗应生效"
    assert len(doc.tables) == 1 and doc.tables[0].cell(1, 0).text == "1", "表格应写入"
    styles = {p.style.name for p in doc.paragraphs}
    assert "List Bullet" in styles and "List Number" in styles, "列表样式应应用"
    assert any("print" in p.text for p in doc.paragraphs), "代码块应写入"


def test_pptx_create_strips_markdown_in_body(tmp_path):
    """回归：pptx body 含 markdown 标记时，不写字面 md 进幻灯片。
    曾把 ## 标题/**加粗**/`代码`/[链接] 原样写入 → 内容显示成 md 文本。"""
    from pptx import Presentation
    out = str(tmp_path / "md.pptx")
    slides = [{
        "title": "背景",
        "body": "## 为什么需要\n- **24小时**自助\n- 覆盖 `挂号` 与 *取号* [详情](https://x.com)\n[图片：排队场景]",
    }]
    r = dc.pptx_create(path=out, slides=slides)
    assert r.startswith("已生成"), r
    prs = Presentation(out)
    body_text = ""
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                body_text += sh.text_frame.text + "\n"
    assert "##" not in body_text, "不应出现字面 ##"
    assert "**" not in body_text, "不应出现字面 **"
    assert "[" not in body_text or "[图片" in body_text, "md 链接应被剥离"
    assert "24小时自助" in body_text.replace("**", ""), "加粗应剥离为纯文本"
    assert "为什么需要" in body_text and "##" not in body_text


def test_pptx_create_cover_image_and_multi_photo(tmp_path):
    """回归：pptx_create 支持封面大图(压暗遮罩叠标题) + 每页多真实图(image 数组，cover裁铺)。"""
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    wide = str(tmp_path / "w.jpg"); Image.new("RGB", (2048, 1000), (30, 80, 140)).save(wide)
    sq = str(tmp_path / "s.png"); Image.new("RGB", (800, 800), (10, 150, 120)).save(sq)
    cover = str(tmp_path / "c.jpg"); Image.new("RGB", (1920, 1080), (10, 40, 80)).save(cover)
    out = str(tmp_path / "e.pptx")
    slides = [
        {"title": "展示", "bullets": ["- 主项"],
         "image": [{"path": wide, "caption": "图A"}, {"path": sq, "caption": "图B"}]},
        {"title": "单图", "image": sq},
    ]
    r = dc.pptx_create(path=out, slides=slides, title="自助挂号机", cover_image=cover,
                       cover_subtitle="智慧医疗")
    assert r.startswith("已生成"), r
    prs = Presentation(out)
    slides_objs = list(prs.slides)
    # 封面页(第1张)应有：背景图(封面) + 压暗遮罩 + 标题文本
    cover_pics = [sh for sh in slides_objs[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert cover_pics, "封面应有背景图"
    cover_txt = "".join(sh.text_frame.text for sh in slides_objs[0].shapes
                        if sh.has_text_frame and sh.text_frame.text)
    assert "自助挂号机" in cover_txt and "智慧医疗" in cover_txt
    # 多图页(第2张)应有 >=2 张图
    pics2 = [sh for sh in slides_objs[1].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics2) >= 2, f"多图页应有 >=2 张真实图，实际 {len(pics2)}"


def test_html_render_produces_png(tmp_path):
    """html_render：把 HTML/CSS 渲染成 PNG（AI 以 HTML 做专业排版的输出通道）。
    需 playwright + 系统 Edge/chromium；缺则跳过。"""
    pytest = __import__("pytest")
    try:
        import playwright  # noqa: F401
    except Exception:
        pytest.skip("未安装 playwright")
    out = str(tmp_path / "d.png")
    html = "<html><body style='margin:0'><div style='width:100vw;height:100vh;background:#0B3D63;color:#fff;display:flex;align-items:center;justify-content:center;font-family:sans-serif'>Hello 自助机</div></body></html>"
    r = dc.html_render(html=html, output=out, width=640, height=360)
    assert r.startswith("已渲染"), r
    assert os.path.exists(out) and os.path.getsize(out) > 0


def test_html_to_ppt_makes_full_bleed_slides(tmp_path):
    """html_to_ppt：多段 HTML 设计 → 每页全幅图的整份 PPT（HTML 专业排版闭环）。"""
    pytest = __import__("pytest")
    try:
        import playwright  # noqa: F401
    except Exception:
        pytest.skip("未安装 playwright")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    page1 = "<div style='width:100vw;height:100vh;background:#0B3D63;color:#fff'>封面</div>"
    page2 = "<div style='width:100vw;height:100vh;background:#fff;color:#123'>内容</div>"
    out = str(tmp_path / "h.pptx")
    r = dc.html_to_ppt(path=out, pages=[page1, page2], width=640, height=360, scale=1)
    assert r.startswith("已生成"), r
    prs = Presentation(out)
    slides = list(prs.slides)
    assert len(slides) == 2, "两页 HTML 应生成两页"
    for s in slides:
        pics = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert pics, "每页应有全幅图"


def test_find_images_search_local_assets(tmp_path):
    """find_images：按关键词/尺寸从本地素材目录检索图片（替代占位选真实图）。"""
    from PIL import Image
    mat = tmp_path / "素材"; (mat / "案例").mkdir(parents=True)
    Image.new("RGB", (400, 300), (200, 30, 30)).save(str(mat / "挂号机_a.jpg"))
    Image.new("RGB", (2048, 1536), (10, 120, 90)).save(str(mat / "挂号机_高清.png"))
    Image.new("RGB", (40, 40), (0, 0, 200)).save(str(mat / "icon.png"))
    Image.new("RGB", (1024, 768), (30, 80, 140)).save(str(mat / "案例" / "医院_现场.jpg"))
    r = dc.find_images(str(mat), keyword="挂号", min_width=300)
    assert "挂号机_高清" in r and "2048x1536" in r, r
    assert "icon" not in r, "过小图应被 min_width 过滤"
    r2 = dc.find_images(str(mat), keyword="医院")
    assert "医院_现场" in r2, "应递归命中子目录"


def test_html_to_pdf_chinese_embedded(tmp_path):
    """html_to_pdf：HTML/CSS 渲染成印刷级 PDF，中文可提取、支持分页。"""
    pytest = __import__("pytest")
    try:
        import playwright  # noqa: F401
    except Exception:
        pytest.skip("未安装 playwright")
    out = str(tmp_path / "d.pdf")
    html = ("<html><head><meta charset='utf-8'><style>@page{size:A4}"
            "body{font-family:'Microsoft YaHei'}</style></head>"
            "<body><h1>自助挂号机系统</h1><p>中文正文可嵌入。</p>"
            "<div style='page-break-before:always'></div><p>第二页</p></body></html>")
    r = dc.html_to_pdf(html=html, output=out, size="A4")
    assert r.startswith("已生成"), r
    assert os.path.exists(out) and os.path.getsize(out) > 0
    import pymupdf
    d = pymupdf.open(out)
    assert d.page_count >= 2, "应正确分页"
    assert "自助挂号机" in d[0].get_text(), "中文应可提取"
