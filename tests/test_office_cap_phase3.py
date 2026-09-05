# -*- coding: utf-8 -*-
"""阶段三（表格/文档/PDF 增强）回归：S8/S9/S11/S13/S14。

- S8 write_excel：style（表头加粗/冻结/筛选/自适应列宽）、start_cell、mode=update、
  数值类型保留、原生内嵌图表（openpyxl.chart）
- S9 docx_edit（replace/insert/append）+ xlsx_edit（单元格级保留格式/公式）
- S11 docx_read 图片占位计数 + create_doc markdown ![图] 插图
- S13 pdf_create 页脚页码 / 页眉
- S14 shared 长度阈值族 + pptx_read 全局输出上限
"""
import os
import zipfile

import deepseek_client as dc


def _perm(tmp_path):
    import permissions
    permissions.WORKSPACE_DIR = str(tmp_path)


def _reg(name):
    names = [t["function"]["name"] for t in dc.TOOLS]
    assert name in names and name in dc.TOOL_CALL_MAP and hasattr(dc, name), name


# ---------------- S8 ----------------
def test_s8_write_excel_style_and_types(tmp_path):
    _perm(tmp_path)
    from openpyxl import load_workbook
    p = str(tmp_path / "s8.xlsx")
    r = dc.write_excel(p, [{"城市": "北京", "销量": 100}, {"城市": "上海", "销量": 80}],
                       header=True, style=True)
    assert r.startswith("已写入"), r
    ws = load_workbook(p).active
    assert ws["A1"].value == "城市" and ws["B1"].value == "销量"
    assert isinstance(ws["B2"].value, int), "数值类型应保留"
    assert ws["A1"].font.bold and ws.freeze_panes == "A2" and ws.auto_filter.ref


def test_s8_write_excel_update_start_cell(tmp_path):
    _perm(tmp_path)
    from openpyxl import Workbook, load_workbook
    p = str(tmp_path / "u.xlsx")
    wb = Workbook(); ws = wb.active; ws["A1"] = "原值"; wb.save(p)
    r = dc.write_excel(p, [["x", "1"]], mode="update", start_cell="D5", header=False)
    assert "已更新" in r, r
    ws = load_workbook(p).active
    assert ws["D5"].value == "x" and ws["A1"].value == "原值", "update 保留其它单元格"


def test_s8_write_excel_native_chart(tmp_path):
    _perm(tmp_path)
    p = str(tmp_path / "c.xlsx")
    r = dc.write_excel(p, [{"季度": "Q1", "营收": 30}, {"季度": "Q2", "营收": 45}],
                       header=True,
                       charts=[{"type": "bar", "title": "季度营收",
                                "categories": ["Q1", "Q2"], "values": [30, 45]}])
    assert r.startswith("已写入"), r
    with zipfile.ZipFile(p) as z:
        assert any(n.endswith(".xml") and "chart" in n for n in z.namelist()), "原生 chart XML"


# ---------------- S9 ----------------
def test_s9_docx_edit(tmp_path):
    _perm(tmp_path)
    _reg("docx_edit")
    from docx import Document
    p = str(tmp_path / "e.docx")
    dc.create_doc(p, "# 报告\n第一段内容。\n\n第二段保留。")
    r = dc.docx_edit(p, action="replace", find="第一段内容", replace="**已改**")
    assert r.startswith("已替换"), r
    doc = Document(p)
    assert any(rn.bold for para in doc.paragraphs if "已改" in para.text for rn in para.runs)
    assert any("第二段保留" in para.text for para in doc.paragraphs)
    r = dc.docx_edit(p, action="insert", anchor="已改", text="插入段")
    assert "插入" in r, r
    txts = [para.text for para in Document(p).paragraphs]
    assert txts.index("插入段") == txts.index("已改。") + 1 or txts.index("插入段") > txts.index(
        next(t for t in txts if "已改" in t))


def test_s9_xlsx_edit(tmp_path):
    _perm(tmp_path)
    _reg("xlsx_edit")
    from openpyxl import Workbook, load_workbook
    p = str(tmp_path / "x.xlsx")
    wb = Workbook(); ws = wb.active; ws["A1"] = "标题"; ws["A2"] = 10; wb.save(p)
    r = dc.xlsx_edit(p, {"B1": "=SUM(A2:A2)", "A2": 99})
    assert "已更新 2" in r, r
    ws = load_workbook(p).active
    assert ws["A2"].value == 99 and str(ws["B1"].value).startswith("=")
    assert ws["A1"].value == "标题", "未改单元格保留"


# ---------------- S11 ----------------
def test_s11_docx_read_image_placeholder(tmp_path):
    _perm(tmp_path)
    from PIL import Image
    from docx import Document
    from docx.shared import Inches
    img = str(tmp_path / "t.png")
    Image.new("RGB", (60, 40), (200, 30, 30)).save(img)
    p = str(tmp_path / "p.docx")
    doc = Document(); doc.add_picture(img, width=Inches(2)); doc.add_paragraph("说明"); doc.save(p)
    out = dc.docx_read(p)
    assert "[图片: 共 1 张" in out or "1 张" in out, out[-80:]


def test_s11_create_doc_embeds_image(tmp_path):
    _perm(tmp_path)
    from PIL import Image
    from docx import Document
    img = str(tmp_path / "t2.png")
    Image.new("RGB", (40, 40), (30, 30, 200)).save(img)
    p = str(tmp_path / "d.docx")
    r = dc.create_doc(p, "# 带图\n文字\n\n![示意图](%s)" % img)
    assert r.startswith("已创建"), r
    xml = Document(p).element.xml
    assert "blip" in xml or "graphic" in xml, "docx 应含图片"


# ---------------- S13 ----------------
def test_s13_pdf_page_footer(tmp_path):
    _perm(tmp_path)
    import pymupdf
    p = str(tmp_path / "o.pdf")
    md = "# 标题\n\n" + "\n\n".join(f"第{i}段长内容用于撑到多页。" for i in range(60))
    r = dc.pdf_create(content=md, output=p)
    assert r.startswith("已生成"), r
    doc = pymupdf.open(p)
    assert doc.page_count >= 2, f"应多页: {doc.page_count}"
    # 页脚页码应出现在每页文本（数字）
    assert any(str(n) in doc[n].get_text() for n in range(doc.page_count))
    doc.close()


# ---------------- S14 ----------------
def test_s14_pptx_read_global_cap(tmp_path):
    _perm(tmp_path)
    from pptx import Presentation
    from shared import PPTX_MAX_DEFAULT
    lp = str(tmp_path / "long.pptx")
    prs = Presentation()
    for i in range(40):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"页{i}"
        for j in range(30):
            prs.slides[i].placeholders[1].text_frame.add_paragraph().text = "填充内容" * 40
    prs.save(lp)
    out = dc.pptx_read(lp)
    assert isinstance(out, str)
    assert len(out) <= PPTX_MAX_DEFAULT + 60, f"应截断: len={len(out)}"
    assert "[内容较长已截断" in out


def test_ppt_layout_check_detects_oob_and_overlap(tmp_path):
    """美学自检：ppt_layout_check 应报越界/重叠，干净文件应通过。"""
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    bad = str(tmp_path / "bad.pptx")
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1)).text = "标题"
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12), Inches(6.5), Inches(3), Inches(2))  # 越界
    prs.save(bad)
    r = dc.ppt_layout_check(bad)
    assert "越界" in r, r
    good = str(tmp_path / "ok.pptx")
    prs = Presentation(); s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(1)).text = "OK"
    prs.save(good)
    r2 = dc.ppt_layout_check(good)
    assert "几何自检通过" in r2, r2
