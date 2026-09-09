# -*- coding: utf-8 -*-
"""阶段一（文档/表格/PPT 能力止血）回归：S1–S5。

覆盖报告《文档表格PPT能力报告_v3.8.5.md》阶段一止血项：
- S1 create_doc 扩展名白名单（pptx/pdf 拒绝，不再静默降级 md）
- S2 write_excel overwrite 覆盖保护（备份 .bak + 重建警告）
- S3 read_csv 编码回退链（GBK 不乱码）
- S4 read_excel 公式无缓存值提示
- S5 read_excel/read_csv 输出统一 markdown 表格 + 单元格 | 转义 + has_header
附加：write_csv 追加模式
"""
import os

import deepseek_client as dc
from openpyxl import Workbook


def _perm(tmp_path):
    import permissions
    permissions.WORKSPACE_DIR = str(tmp_path)


def test_s1_create_doc_rejects_pptx_pdf(tmp_path):
    _perm(tmp_path)
    for ext in (".pptx", ".pdf", ".ppt", ".txt"):
        r = dc.create_doc(str(tmp_path / f"f{ext}"), "# hi")
        assert r.startswith("错误"), f"{ext} 应被拒绝"
        assert ext.lstrip(".") in r
    # md 仍可用
    ok = dc.create_doc(str(tmp_path / "ok.md"), "# 标题")
    assert ok.startswith("已创建")


def test_s2_write_excel_overwrite_backs_up(tmp_path):
    _perm(tmp_path)
    p = str(tmp_path / "x.xlsx")
    wb = Workbook(); ws = wb.active; ws.title = "S1"; ws.append(["x"])
    wb.create_sheet("keepme").append(["secret"]); wb.save(p)
    r = dc.write_excel(p, [["a", "1"]], sheet="new")
    assert "已备份" in r, "overwrite 已存在文件应提示已备份"
    assert os.path.exists(p + ".bak"), "应生成 .bak"
    from openpyxl import load_workbook
    assert load_workbook(p).sheetnames == ["new"], "overwrite 重建仅含新 sheet"


def test_s2b_write_excel_array_rows_with_headers_preserves_data(tmp_path):
    """回归：write_excel「数组行 + 显式 headers」必须真实写入数据，不得假成功丢数据。

    原 Bug：数组行 + headers 时，cols 被 headers 填成非 None → 误走 dict 分支 →
    数组行被当非 dict 填全空串 → 返回"已写入 N 行"但底层数据行全空（静默丢数据）。
    此用例直接调用 + openpyxl 回读，锁死"数据行必须非空"。
    """
    _perm(tmp_path)
    from openpyxl import load_workbook

    p = str(tmp_path / "arr.xlsx")
    data = [["1月", 120, 80, 40], ["2月", 150, 95, 55]]
    r = dc.write_excel(p, data=data, headers=["月份", "销售", "成本", "利润"])
    assert "已写入" in r
    ws = load_workbook(p).active
    # 表头
    assert [c.value for c in ws[1]] == ["月份", "销售", "成本", "利润"]
    # 数据行必须真实存在（防假成功丢数据）
    assert [c.value for c in ws[2]] == ["1月", 120, 80, 40], f"第2行丢失: {[c.value for c in ws[2]]}"
    assert [c.value for c in ws[3]] == ["2月", 150, 95, 55], f"第3行丢失: {[c.value for c in ws[3]]}"


def test_s2c_write_excel_object_rows_with_headers(tmp_path):
    """对象数组 + headers 也必须正常（原有路径不回归）。"""
    _perm(tmp_path)
    from openpyxl import load_workbook

    p = str(tmp_path / "obj.xlsx")
    r = dc.write_excel(p, data=[{"月份": "1月", "销售": 120}], headers=["月份", "销售"])
    assert "已写入" in r
    ws = load_workbook(p).active
    assert [c.value for c in ws[1]] == ["月份", "销售"]
    assert [c.value for c in ws[2]] == ["1月", 120]


def test_s3_read_csv_gbk_fallback(tmp_path):
    _perm(tmp_path)
    p = str(tmp_path / "gbk.csv")
    with open(p, "w", encoding="gbk", newline="") as f:
        f.write("姓名,城市\n张三,杭州\n")
    out = dc.read_csv(p)
    assert "张三" in out and "杭州" in out, "GBK 中文不应乱码"
    assert "gb18030" in out, "应标注实际探测编码"


def test_s4_read_excel_formula_hint(tmp_path):
    _perm(tmp_path)
    p = str(tmp_path / "f.xlsx")
    wb = Workbook(); ws = wb.active
    ws.append(["a", "b"]); ws.append([1, 2]); ws["B2"] = "=A2*2"
    wb.save(p)
    out = dc.read_excel(p)
    assert "含公式" in out and "B2" in out, "公式列无缓存值时应给出提示"


def test_s5_table_outputs_markdown_and_escape_pipe(tmp_path):
    _perm(tmp_path)
    # CSV 单元格含 | 不串列
    cp = str(tmp_path / "p.csv")
    with open(cp, "w", encoding="utf-8-sig", newline="") as f:
        f.write("A,B\nx|y,z\n")
    out = dc.read_csv(cp)
    assert out.lstrip().startswith("|"), "应输出 markdown 表格"
    assert "x\\|y" in out, "单元格内 | 应转义"
    assert out.count("x\\|y") == 1
    # has_header=false 补占位表头
    assert "col_1" in dc.read_csv(cp, has_header=False)
    # Excel 同契约
    xp = str(tmp_path / "t.xlsx")
    wb = Workbook(); ws = wb.active; ws.append(["a", "1"]); wb.save(xp)
    assert dc.read_excel(xp).lstrip().startswith("|")
    assert "col_1" in dc.read_excel(xp, has_header=False)


def test_write_csv_append_mode(tmp_path):
    _perm(tmp_path)
    p = str(tmp_path / "w.csv")
    dc.write_csv(p, [["a", "1"]], headers="k,v")
    r = dc.write_csv(p, [["c", "3"]], mode="append")
    assert "追加" in r
    content = open(p, encoding="utf-8-sig").read()
    assert content.count("k,v") == 1, "追加不应重复表头"
    assert "c,3" in content
